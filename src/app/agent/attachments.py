"""Agent 附件摄取：按类型提取可理解上下文，不执行文件内容。"""

import base64
import hashlib
import io
import json
import mimetypes
import subprocess
import zipfile
from pathlib import Path
from typing import cast

import bencodepy
from charset_normalizer import from_bytes
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from rarfile import RarFile

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 50000
MAX_ARCHIVE_TEXT_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 50

_TEXT_SUFFIXES = {
    ".txt",
    ".log",
    ".md",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".conf",
    ".csv",
    ".xml",
    ".py",
    ".js",
    ".ts",
    ".sh",
}


def _decode_text(data: bytes) -> str:
    result = from_bytes(data).best()
    if result is None:
        return data.decode("utf-8", errors="replace")
    return str(result)


def _truncate(text: str, suffix: str = "") -> tuple[str, bool]:
    if len(text) <= MAX_TEXT_CHARS:
        return text, False
    return text[:MAX_TEXT_CHARS] + f"\n\n[内容已截断{suffix}]", True


def _torrent_text(data: bytes) -> str:
    decoded = cast(dict, bencodepy.decode(data))
    info = cast(dict, decoded.get(b"info") or {})
    name = (info.get(b"name") or b"").decode("utf-8", errors="replace")
    files = list(info.get(b"files") or [])
    lines = [f"种子名称：{name}"]
    if files:
        lines.append(f"文件数量：{len(files)}")
        for raw_item in files[:100]:
            item = cast(dict, raw_item)
            path = b"/".join(item.get(b"path") or []).decode("utf-8", errors="replace")
            lines.append(f"- {path} ({int(item.get(b'length') or 0)} bytes)")
    elif info.get(b"length"):
        lines.append(f"文件大小：{int(info[b'length'])} bytes")
    announce = decoded.get(b"announce")
    if announce:
        lines.append(f"Tracker：{announce.decode('utf-8', errors='replace')}")
    info_hash = hashlib.sha1(bencodepy.encode(info)).hexdigest()
    lines.append(f"Info Hash：{info_hash}")
    return "\n".join(lines)


def _zip_text(data: bytes) -> str:
    lines = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()[:MAX_ARCHIVE_ENTRIES]
        lines.append(f"压缩包文件数：{len(archive.infolist())}")
        budget = MAX_ARCHIVE_TEXT_BYTES
        for info in infos:
            lines.append(f"- {info.filename} ({info.file_size} bytes)")
            suffix = Path(info.filename).suffix.lower()
            if info.is_dir() or suffix not in _TEXT_SUFFIXES or info.file_size > budget:
                continue
            content = archive.read(info)[: min(info.file_size, 200_000)]
            budget -= len(content)
            lines.append(f"\n[{info.filename}]\n{_decode_text(content)[:10000]}")
            if budget <= 0:
                lines.append("[压缩包文本预算已用完]")
                break
    return "\n".join(lines)


def _rar_text(data: bytes) -> str:
    lines = []
    with RarFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()[:MAX_ARCHIVE_ENTRIES]
        lines.append(f"压缩包文件数：{len(archive.infolist())}")
        budget = MAX_ARCHIVE_TEXT_BYTES
        for info in infos:
            lines.append(f"- {info.filename} ({info.file_size} bytes)")
            suffix = Path(info.filename).suffix.lower()
            if info.is_dir() or suffix not in _TEXT_SUFFIXES or info.file_size > budget:
                continue
            content = archive.read(info)[: min(info.file_size, 200_000)]
            budget -= len(content)
            lines.append(f"\n[{info.filename}]\n{_decode_text(content)[:10000]}")
    return "\n".join(lines)


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for index, page in enumerate(reader.pages[:80], 1):
        parts.append(f"[第 {index} 页]\n{page.extract_text() or ''}")
    return "\n\n".join(parts)


def _docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets[:10]:
        parts.append(f"[工作表 {sheet.title}]")
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= 300:
                break
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                parts.append(" | ".join(values))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    parts = []
    for index, slide in enumerate(list(presentation.slides)[:80], 1):
        texts = [text for shape in slide.shapes if (text := str(getattr(shape, "text", "") or "")).strip()]
        parts.append(f"[幻灯片 {index}]\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _image_text(data: bytes) -> str:
    image = Image.open(io.BytesIO(data))
    return (
        f"图片格式：{image.format}\n尺寸：{image.width} x {image.height}\n"
        f"模式：{image.mode}\n说明：当前模型未配置视觉识别，仅提取图片元数据。"
    )


def _media_text(data: bytes, suffix: str) -> str:
    try:
        result = subprocess.run(
            [
                "ffprobe",
                "-v",
                "error",
                "-show_format",
                "-show_streams",
                "-of",
                "json",
                "pipe:0",
            ],
            input=data,
            capture_output=True,
            timeout=15,
            check=False,
        )
        payload = json.loads(result.stdout or b"{}")
        streams = payload.get("streams") or []
        lines = [f"媒体容器：{(payload.get('format') or {}).get('format_name', suffix)}"]
        for stream in streams[:10]:
            lines.append(
                f"- {stream.get('codec_type')}: {stream.get('codec_name')} "
                f"{stream.get('width', '')}x{stream.get('height', '')} "
                f"duration={stream.get('duration', '')}"
            )
        return "\n".join(lines)
    except Exception as error:  # noqa: BLE001
        return f"音视频文件，ffprobe 解析失败：{error}"


def extract_attachment(name: str, data: bytes) -> dict:
    """提取附件上下文；返回统一结构，不执行或落盘原始文件。"""
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("附件不能超过 25 MB")
    suffix = Path(name or "").suffix.lower()
    kind = "binary"
    try:
        if suffix in _TEXT_SUFFIXES:
            kind, text = "text", _decode_text(data)
        elif suffix == ".pdf":
            kind, text = "pdf", _pdf_text(data)
        elif suffix == ".docx":
            kind, text = "docx", _docx_text(data)
        elif suffix == ".xlsx":
            kind, text = "xlsx", _xlsx_text(data)
        elif suffix == ".pptx":
            kind, text = "pptx", _pptx_text(data)
        elif suffix == ".torrent":
            kind, text = "torrent", _torrent_text(data)
        elif suffix == ".zip":
            kind, text = "archive", _zip_text(data)
        elif suffix == ".rar":
            kind, text = "archive", _rar_text(data)
        elif suffix in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}:
            kind, text = "image", _image_text(data)
        elif suffix in {".mp4", ".mkv", ".avi", ".mov", ".webm", ".mp3", ".flac", ".wav", ".m4a"}:
            kind, text = "media", _media_text(data, suffix)
        else:
            mime = mimetypes.guess_type(name or "")[0] or "application/octet-stream"
            text = (
                f"二进制文件\nMIME：{mime}\n大小：{len(data)} bytes\n"
                f"SHA256：{hashlib.sha256(data).hexdigest()}\n"
                f"文件头：{data[:64].hex()}"
            )
    except Exception as error:  # noqa: BLE001
        raise ValueError(f"附件解析失败：{error}") from error
    text, truncated = _truncate(text, f"，原始大小 {len(data)} bytes")
    result = {
        "name": name,
        "kind": kind,
        "text": text,
        "chars": len(text),
        "truncated": truncated,
        "size": len(data),
    }
    if kind == "image":
        result["image_data"] = base64.b64encode(data).decode()
        result["media_type"] = mimetypes.guess_type(name or "")[0] or "image/png"
    return result
